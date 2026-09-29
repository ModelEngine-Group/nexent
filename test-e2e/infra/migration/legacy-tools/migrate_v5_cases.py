"""One-time, behavior-preserving conversion of the canonical V5 case catalog.

The converter never edits the source suite. It writes only new migration-owned
Feature and Case documents; automation and manifest migration are a later gate.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import defaultdict
from pathlib import Path

import yaml
from openpyxl import load_workbook


CASE_GROUPS = ("main_cases", "journeys", "special_cases")
GROUP_SHEETS = {
    "main_cases": "02_测试用例基线",
    "journeys": "03_Playwright_Journey",
    "special_cases": "04_安全性能可靠性",
}
NUMBERED_LINE = re.compile(r"^\s*\d+\s*[.、)）]\s*(.*)$")
LEGACY_PRODUCT_SHA = re.compile(r"\b628d3e72(?:a97dd5e8773aea62fbe8bb98127ab917)?\b")
SKIPPED_BY_POLICY = {"PW-AUTH-03", "PW-AUTH-04"}
RETIRED_BY_PRODUCT = {"AGT-072", "PW-HITL-01"}
NO_FORMAL_IMPLEMENTATION = {
    "PW-A2A-DISCOVERY-01", "PW-A2A-PUBLISH-01",
    "PW-KB-AIDP-01", "PW-MCP-LOCALIMG-01", "PW-QUOTA-LOCAL-01",
    "CTR-034", "CTR-035", "CTR-039", "CTR-040", "CTR-041", "CTR-042", "CTR-043",
    "API-092", "API-093", "CTR-031", "CTR-032", "CTR-033", "CTR-036",
    "CTR-037", "CTR-038", "API-094", "API-095",
}


class LiteralDumper(yaml.SafeDumper):
    """Keep the original multi-line V5 steps readable in Git diffs."""

    def ignore_aliases(self, data):
        return True


def _represent_string(dumper: yaml.Dumper, value: str):
    return dumper.represent_scalar("tag:yaml.org,2002:str", value, style="|" if "\n" in value else None)


LiteralDumper.add_representer(str, _represent_string)


def text(value: object) -> str:
    # The old workbook pinned one historical develop commit in generic setup
    # prose. Formal cases now travel with the product commit that owns them.
    return LEGACY_PRODUCT_SHA.sub("当前批次锁定的产品提交", str(value or "")).strip()


def clauses(value: object) -> list[str]:
    """Split numbered lines without dropping their continuation text."""
    source = text(value)
    if not source:
        return []
    items: list[str] = []
    for line in source.splitlines():
        match = NUMBERED_LINE.match(line)
        if match:
            items.append(match.group(1).strip())
        elif items:
            items[-1] += "\n" + line
        elif line.strip():
            items.append(line.strip())
    return [item.strip() for item in items if item.strip()]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def workbook_cases(workbook) -> dict[str, list[dict]]:
    """Read canonical rows directly; the old JSON catalog is only a cross-check."""
    result: dict[str, list[dict]] = {}
    for group, sheet_name in GROUP_SHEETS.items():
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows)
        group_rows: list[dict] = []
        for row_number, values in enumerate(rows, start=2):
            if not values or not values[0]:
                continue
            record = {header: value for header, value in zip(headers, values) if header is not None}
            identifier = str(values[0]).strip()
            stage_match = re.match(r"^(D[1-5])\b", str(record.get("每日CI阶段") or ""))
            if not stage_match:
                raise ValueError(f"{sheet_name}:{row_number}: no D1-D5 stage")
            record.update(id=identifier, stage=stage_match.group(1), _workbook_row=row_number)
            group_rows.append(record)
        result[group] = group_rows
    return result


def risk_type(case: dict) -> str:
    label = text(case.get("专项类型") or case.get("测试方式")).lower()
    identifier = text(case.get("id"))
    if "security" in label or "安全" in label or identifier.startswith("SEC-"):
        return "SECURITY"
    if "performance" in label or "性能" in label or identifier.startswith("PERF-"):
        return "PERFORMANCE"
    if "deployment" in label or "部署" in label or identifier.startswith("DEP-"):
        return "DEPLOYMENT"
    return "RELIABILITY"


def legacy_fields(case: dict, group: str) -> tuple[str, str, str, str, str, str]:
    if group == "main_cases":
        return (
            text(case.get("用例标题")), text(case.get("前置条件")),
            text(case.get("详细测试步骤")), text(case.get("预期结果/断言")),
            text(case.get("本地资产/配置")), text(case.get("失败证据/产物")),
        )
    if group == "journeys":
        return (
            text(case.get("场景名称")), text(case.get("账号与前置条件")),
            text(case.get("UI详细步骤（Playwright）")), text(case.get("关键断言")),
            text(case.get("本地资产/配置")), text(case.get("失败证据/产物")),
        )
    return (
        text(case.get("测试项")), text(case.get("本地资产/配置")),
        text(case.get("详细测试方法")), text(case.get("通过标准")),
        text(case.get("本地资产/配置")), text(case.get("失败证据/产物")),
    )


def convert(source_root: Path) -> tuple[dict, dict[str, dict], dict]:
    source_root = source_root.resolve()
    workbook_path = source_root / "cases/Nexent_Develop_每日真实资产测试基线_v5.xlsx"
    catalog = json.loads((source_root / "auto_test/case_catalog.json").read_text(encoding="utf-8"))
    workbook_hash = sha256(workbook_path)
    if catalog.get("source_sha256", "").lower() != workbook_hash:
        raise ValueError("The V5 catalog does not match the canonical workbook")
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    source_cases = workbook_cases(workbook)
    stale_catalog = []
    for group in CASE_GROUPS:
        indexed = {entry["id"]: entry for entry in catalog[group]}
        if set(indexed) != {entry["id"] for entry in source_cases[group]}:
            raise ValueError(f"Workbook/catalog Case IDs differ in {group}")
        for current in source_cases[group]:
            old = indexed[current["id"]]
            if any(text(current.get(key)) != text(old.get(key))
                   for key in current if key not in {"id", "stage", "_workbook_row"}):
                stale_catalog.append(current["id"])
    features: dict[str, dict] = {}
    if "01_代码功能清单" not in workbook.sheetnames:
        raise ValueError("Canonical V5 feature sheet 01_代码功能清单 is missing")
    for row in list(workbook["01_代码功能清单"].values)[1:]:
        if not row or not row[0]:
            continue
        identifier = text(row[0])
        name = text(row[3])
        description = f"{text(row[1])} / {text(row[2])}: {name}"
        features[identifier] = {
            "feature_id": identifier,
            "name": name,
            "status": "active",
            "description": description,
            "business_rules": [{"rule_id": f"{identifier}-RULE-01", "description": name}],
            "required_test_stages": [],
            "source": {"notes": "Migrated from V5 01_代码功能清单 without changing the functional description."},
        }

    stage_cases: dict[str, list[dict]] = defaultdict(list)
    report = {"source_sha256": workbook_hash, "stale_catalog_case_ids": stale_catalog,
              "provisional_feature_links": [], "blocked_cases": [],
              "policy_skipped_cases": [], "retired_cases": [], "normalized_legacy_commit_pin_cases": []}
    for group in CASE_GROUPS:
        for original in source_cases[group]:
            row_number = original["_workbook_row"]
            identifier = text(original["id"])
            stage = text(original["stage"])
            if LEGACY_PRODUCT_SHA.search(json.dumps(original, ensure_ascii=False)):
                report["normalized_legacy_commit_pin_cases"].append(identifier)
            title, before, actions, expected, assets, evidence = legacy_fields(original, group)
            feature_id = text(original.get("功能ID")) if group == "main_cases" else ""
            if feature_id not in features:
                feature_id = f"V5-{identifier}"
                features[feature_id] = {
                    "feature_id": feature_id,
                    "name": title or identifier,
                    "status": "active",
                    "description": text(original.get("业务目标") or original.get("风险目标") or title or identifier),
                    "business_rules": [{"rule_id": f"{feature_id}-RULE-01", "description": title or identifier}],
                    "required_test_stages": [],
                    "source": {"notes": "Provisional migration owner: the V5 row has no stable Feature ID. Do not infer an F-ID."},
                }
                report["provisional_feature_links"].append(identifier)
            if identifier not in RETIRED_BY_PRODUCT:
                features[feature_id]["required_test_stages"] = sorted(set(features[feature_id]["required_test_stages"]) | {stage})

            preconditions = clauses(before)
            expected_results = clauses(expected) or [title or identifier]
            action_items = clauses(actions) or [title or identifier]
            steps = [
                {"order": index, "action": action, "expected": expected_results[min(index - 1, len(expected_results) - 1)]}
                for index, action in enumerate(action_items, start=1)
            ]
            if identifier in RETIRED_BY_PRODUCT:
                status, automation = "retired", "not_applicable"
                report["retired_cases"].append(identifier)
            elif identifier in SKIPPED_BY_POLICY:
                status, automation = "skipped_by_policy", "not_applicable"
                report["policy_skipped_cases"].append(identifier)
            elif identifier in NO_FORMAL_IMPLEMENTATION:
                status, automation = "blocked", "automated"
                report["blocked_cases"].append(identifier)
            else:
                status, automation = "active", "automated"
            if stage == "D1":
                case_type = text(original.get("测试方式"))
            elif stage == "D2":
                case_type = text(original.get("测试方式"))
            elif stage == "D3":
                case_type = "AGENT-IT" if original.get("测试方式") == "AGENT-IT" else "INTEGRATION"
            elif stage == "D4":
                case_type = "E2E"
            else:
                case_type = risk_type(original)
            case = {
                "case_id": identifier,
                "feature_id": feature_id,
                "business_rule_ids": [f"{feature_id}-RULE-01"],
                "title": title or identifier,
                "type": case_type,
                "priority": text(original.get("优先级")) or "P1",
                "status": status,
                "automation": automation,
                "objective": text(original.get("业务目标") or original.get("风险目标") or title or identifier),
                "preconditions": preconditions,
                "test_data": {"legacy_asset_declaration": assets} if assets else {},
                "steps": steps,
                "expected_results": expected_results,
                "forbidden_side_effects": [],
                "source": {"legacy_workbook_sha256": workbook_hash, "legacy_sheet": {
                    "main_cases": "02_测试用例基线", "journeys": "03_Playwright_Journey",
                    "special_cases": "04_安全性能可靠性",
                }[group], "legacy_row": row_number},
            }
            if stage == "D1":
                case.update(unit_boundary=text(original.get("子模块") or title or identifier), inputs=[assets] if assets else [], assertions=expected_results)
            elif stage == "D2":
                case["contract"] = {
                    "protocol": "HTTP API" if case_type == "API-IT" else "V5 protocol contract",
                    "request_expectations": action_items,
                    "response_expectations": expected_results,
                }
            elif stage == "D3":
                case.update(integration_boundary=[text(original.get("外部依赖策略") or original.get("子模块") or title or identifier)], observations=expected_results, cleanup=[text(original.get("失败证据/产物") or "V5 case-owned cleanup")])
            elif stage == "D4":
                case.update(actor="Nexent browser user", journey=action_items, final_business_result=expected_results, failure_evidence=clauses(evidence))
            else:
                risk = risk_type(original)
                case.update(risk_type=risk, risk=text(original.get("风险目标") or title or identifier), condition=preconditions or [title or identifier], metrics=expected_results, thresholds=expected_results, recovery_expectations=expected_results)
            stage_cases[stage].append(case)

    feature_document = {"schema_version": "1.0", "module": {"id": "V5-BASELINE", "name": "Migrated V5 baseline"}, "features": list(features.values())}
    case_documents = {
        stage: {"schema_version": "1.0", "stage": stage, "module": "V5-BASELINE", "cases": cases}
        for stage, cases in sorted(stage_cases.items())
    }
    report.update({"feature_count": len(features), "case_count": sum(len(cases) for cases in stage_cases.values())})
    return feature_document, case_documents, report


def dump_yaml(data: dict) -> str:
    return yaml.dump(data, Dumper=LiteralDumper, allow_unicode=True, sort_keys=False, width=120)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-root", type=Path, required=True)
    parser.add_argument("--repo-root", type=Path, default=Path(__file__).resolve().parents[2])
    parser.add_argument("--apply", action="store_true", help="Write new migration-owned Feature and Case documents")
    parser.add_argument("--replace", action="store_true", help="Replace only previously generated V5 migration documents")
    args = parser.parse_args()
    feature_doc, case_docs, report = convert(args.source_root)
    targets = [(args.repo_root / "test/features/v5-baseline.yaml", feature_doc)] + [
        (args.repo_root / f"test/cases/{stage.lower()}/v5-baseline.yaml", document)
        for stage, document in case_docs.items()
    ]
    if args.apply:
        existing = [str(path) for path, _ in targets if path.exists()]
        if existing and not args.replace:
            raise SystemExit("Migration-owned output already exists; refusing overwrite: " + ", ".join(existing))
        if args.replace:
            for path, _ in targets:
                if not path.exists():
                    continue
                previous = yaml.safe_load(path.read_text(encoding="utf-8"))
                module = previous.get("module")
                owner = module.get("id") if isinstance(module, dict) else module
                if owner != "V5-BASELINE":
                    raise SystemExit(f"Refusing to replace a non-migration document: {path}")
        for path, document in targets:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(dump_yaml(document), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
