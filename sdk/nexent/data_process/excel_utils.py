from io import BytesIO


def load_excel_workbook(file_data: bytes):
    """Load XLS/XLSX bytes into a workbook usable by the shared Excel pipeline."""
    import openpyxl

    # Split XLS parts are serialized as XLSX but retain their source filename.
    if not file_data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return openpyxl.load_workbook(BytesIO(file_data))

    import xlrd

    source = xlrd.open_workbook(file_contents=file_data, formatting_info=True)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    try:
        for source_sheet in source.sheets():
            sheet = workbook.create_sheet(source_sheet.name)
            for row_index in range(source_sheet.nrows):
                for column_index in range(source_sheet.row_len(row_index)):
                    source_cell = source_sheet.cell(row_index, column_index)
                    if source_cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                        continue
                    value = source_cell.value
                    if source_cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(value, source.datemode)
                    elif source_cell.ctype == xlrd.XL_CELL_BOOLEAN:
                        value = bool(value)
                    elif source_cell.ctype == xlrd.XL_CELL_ERROR:
                        value = xlrd.error_text_from_code.get(value, "#VALUE!")
                    cell = sheet.cell(row=row_index + 1, column=column_index + 1, value=value)
                    if source_cell.ctype == xlrd.XL_CELL_TEXT:
                        # Preserve literal strings starting with '=' across XLSX serialization.
                        cell.data_type = "s"

            for row_start, row_end, col_start, col_end in source_sheet.merged_cells:
                sheet.merge_cells(
                    start_row=row_start + 1, end_row=row_end,
                    start_column=col_start + 1, end_column=col_end,
                )
        return workbook
    finally:
        source.release_resources()
