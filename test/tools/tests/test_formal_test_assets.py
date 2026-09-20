"""Integration tests for the formal test asset schemas and tools."""

from __future__ import annotations

import json
import shutil
import sys
from pathlib import Path

import yaml
from openpyxl import load_workbook


TOOLS_DIR = Path(__file__).resolve().parents[1]
if str(TOOLS_DIR) not in sys.path:
    sys.path.insert(0, str(TOOLS_DIR))

import generate_excel  # noqa: E402
import validate_test_assets  # noqa: E402
from test_asset_lib import case_contract_hash, implementation_hash  # noqa: E402


def _write_yaml(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(yaml.safe_dump(value, allow_unicode=True, sort_keys=False), encoding="utf-8")


def _common_case(stage: str) -> dict:
    case = {
        "case_id": f"FEATURE-{stage}-001",
        "feature_id": "FEATURE-001",
        "business_rule_ids": ["FEATURE-RULE-001"],
        "title": f"Prove the {stage} behavior",
        "type": {"D1": "BE-UT", "D2": "API-IT", "D3": "INTEGRATION", "D4": "E2E", "D5": "RELIABILITY"}[stage],
        "priority": "P1",
        "status": "active" if stage == "D1" else "manual",
        "automation": "automated" if stage == "D1" else "manual",
        "objective": "Prove the requirement contract at the selected boundary.",
        "preconditions": ["Use the isolated test profile."],
        "test_data": {"profile": "mock"},
        "steps": [{"order": 1, "action": "Exercise the contract.", "expected": "The boundary returns the specified result."}],
        "expected_results": ["The observable result matches the requirement."],
        "forbidden_side_effects": ["No unrelated state changes."],
    }
    if stage == "D1":
        case.update(unit_boundary="service.calculate", inputs=[{"value": 1}], assertions=["The result equals 1."])
    elif stage == "D2":
        case["contract"] = {"protocol": "HTTP", "method": "POST", "endpoint": "/api/example", "request_expectations": ["The body is validated."], "response_expectations": ["The response is stable."]}
    elif stage == "D3":
        case.update(integration_boundary=["service to model-provider mock"], observations=["The provider request is observable."], cleanup=["Delete generated records."])
    elif stage == "D4":
        case.update(actor="Authenticated user", journey=["Open the feature.", "Complete the workflow."], final_business_result=["The completed result is visible."], failure_evidence=["Retain screenshot and trace on failure."])
    else:
        case.update(risk_type="RELIABILITY", risk="A provider timeout could leave partial state.", condition=["Inject a timeout."], metrics=["Partial record count."], thresholds=["Partial record count equals zero."], recovery_expectations=["A retry can complete successfully."])
    return case


def _build_repository(tmp_path: Path, source_root: Path) -> Path:
    root = tmp_path / "repository"
    (root / ".git").mkdir(parents=True)
    shutil.copytree(source_root / "test/schemas", root / "test/schemas")
    _write_yaml(
        root / "test/features/core.yaml",
        {
            "schema_version": "1.0",
            "module": {"id": "MODULE-001", "name": "Core"},
            "features": [{"feature_id": "FEATURE-001", "name": "Example feature", "status": "active", "description": "A complete formal test asset example.", "business_rules": [{"rule_id": "FEATURE-RULE-001", "description": "The feature returns a stable result."}], "required_test_stages": ["D1", "D2", "D3", "D4", "D5"]}],
        },
    )
    cases = {}
    for stage in ("D1", "D2", "D3", "D4", "D5"):
        case = _common_case(stage)
        cases[stage] = case
        _write_yaml(root / f"test/cases/{stage.lower()}/core.yaml", {"schema_version": "1.0", "stage": stage, "module": "MODULE-001", "cases": [case]})
    _write_yaml(
        root / "test/changes/requirements/requirement.yaml",
        {
            "schema_version": "1.0",
            "changes": [{"change_id": "CHANGE-001", "change_type": "requirement", "title": "Add the example feature", "description": "Define formal coverage.", "requirement_id": "REQ-001", "affected_features": {"existing": [], "added": ["FEATURE-001"], "modified": [], "retired": []}, "affected_cases": {"existing": [], "added": [f"FEATURE-D{number}-001" for number in range(1, 6)], "modified": [], "retired": []}}],
        },
    )
    script = root / "test/automation/d1/test_feature.py"
    script.parent.mkdir(parents=True)
    script.write_text("def test_feature_d1_001():\n    assert 1 == 1\n", encoding="utf-8")
    implementations = [{"framework": "pytest", "file": "test/automation/d1/test_feature.py", "selector": "test_feature_d1_001", "profiles": ["mock"]}]
    _write_yaml(
        root / "test/manifests/d1-d5.yaml",
        {
            "schema_version": "1.0",
            "cases": [{"case_id": "FEATURE-D1-001", "stage": "D1", "status": "implemented", "contract_hash": case_contract_hash("D1", cases["D1"]), "implementation_hash": implementation_hash(root, implementations), "implementation": implementations, "required_mock_services": []}],
        },
    )
    return root


def test_unified_validation_and_excel_generation(tmp_path: Path) -> None:
    source_root = Path(__file__).resolve().parents[3]
    root = _build_repository(tmp_path, source_root)

    output = root / "test/generated/Nexent_测试基线.xlsx"
    generate_excel.write_workbook(root, output)
    assert validate_test_assets.validate(root) == []
    workbook = load_workbook(output)
    assert workbook.sheetnames == generate_excel.SHEET_NAMES
    assert workbook["02_D1"]["A2"].value == "FEATURE-D1-001"
    assert workbook["06_D5"]["A2"].value == "FEATURE-D5-001"
    headers = {cell.value: cell.column for cell in workbook["02_D1"][1]}
    assert workbook["02_D1"].cell(row=2, column=headers["自动化实现状态"]).value == "implemented"
    assert workbook["02_D1"].cell(row=2, column=headers["脚本路径"]).value == "test/automation/d1/test_feature.py"
    assert workbook["02_D1"].cell(row=2, column=headers["Manifest校验状态"]).value == "校验通过"
    assert workbook["02_D1"].column_dimensions["Z"].hidden is True
    assert workbook["02_D1"].column_dimensions["AA"].hidden is True
    second_output = root / "test/generated/second.xlsx"
    generate_excel.write_workbook(root, second_output)
    assert generate_excel.workbook_snapshot(output) == generate_excel.workbook_snapshot(second_output)

    workbook["02_D1"]["A2"] = "MANUAL-EDIT"
    workbook.save(output)
    messages = [issue.message for issue in validate_test_assets.validate(root)]
    assert "Generated Excel view is stale or manually modified" in messages


def test_validation_rejects_local_sql_path(tmp_path: Path) -> None:
    source_root = Path(__file__).resolve().parents[3]
    root = _build_repository(tmp_path, source_root)
    case_path = root / "test/cases/d1/core.yaml"
    document = yaml.safe_load(case_path.read_text(encoding="utf-8"))
    document["cases"][0]["test_data"] = {"fixture": "/home/developer/query.sql"}
    _write_yaml(case_path, document)

    messages = [issue.message for issue in validate_test_assets.validate(root)]
    assert "Developer- or runner-local absolute paths are not allowed" in messages
    assert "Business tests must not depend on SQL file paths" in messages


def test_design_phase_does_not_require_manifest(tmp_path: Path) -> None:
    source_root = Path(__file__).resolve().parents[3]
    root = _build_repository(tmp_path, source_root)
    shutil.rmtree(root / "test/manifests")

    assert validate_test_assets.validate(root, phase="design", generate_excel_view=True) == []


def test_change_type_requires_matching_directory(tmp_path: Path) -> None:
    source_root = Path(__file__).resolve().parents[3]
    root = _build_repository(tmp_path, source_root)
    source = root / "test/changes/requirements/requirement.yaml"
    destination = root / "test/changes/requirement.yaml"
    source.replace(destination)

    messages = [issue.message for issue in validate_test_assets.validate(root, phase="design", generate_excel_view=True)]
    assert "Change type requirement must be stored under test/changes/requirements/, not test/changes root" in messages


def test_implementation_hash_detects_script_drift(tmp_path: Path) -> None:
    source_root = Path(__file__).resolve().parents[3]
    root = _build_repository(tmp_path, source_root)
    generate_excel.write_workbook(root, root / "test/generated/Nexent_测试基线.xlsx")
    script = root / "test/automation/d1/test_feature.py"
    script.write_text("def test_feature_d1_001():\n    assert 2 == 2\n", encoding="utf-8")

    messages = [issue.message for issue in validate_test_assets.validate(root)]
    assert any(message.startswith("Implementation hash must be sha256:") for message in messages)


def test_schemas_are_valid_json(tmp_path: Path) -> None:
    source_root = Path(__file__).resolve().parents[3]
    for schema_path in sorted((source_root / "test/schemas").glob("*.json")):
        assert isinstance(json.loads(schema_path.read_text(encoding="utf-8")), dict)
