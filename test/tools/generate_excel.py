"""Generate the deterministic Excel view of formal Nexent test assets."""

from __future__ import annotations

import argparse
import json
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from test_asset_lib import (
    case_contract_hash,
    case_index,
    discover_documents,
    feature_index,
    implementation_hash,
    manifest_index,
    repository_root,
)


SHEET_NAMES = ["00_说明", "01_功能清单", "02_D1", "03_D2", "04_D3", "05_D4", "06_D5"]
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
SECTION_FILL = PatternFill("solid", fgColor="D9EAF7")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(bottom=Side(style="thin", color="B8C2CC"))
WRAP = Alignment(vertical="top", wrap_text=True)
FIXED_TIME = datetime(2000, 1, 1, tzinfo=timezone.utc)


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return str(value)


def _steps(case: dict[str, Any]) -> str:
    return "\n".join(f"{step['order']}. {step['action']} => {step['expected']}" for step in case.get("steps", []))


def _configure_sheet(
    sheet,
    headers: Iterable[str],
    rows: Iterable[Iterable[Any]],
    hidden_headers: set[str] | None = None,
) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for row in rows:
        sheet.append([_text(value) for value in row])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = WRAP
            cell.border = THIN_BORDER
    if sheet.max_column:
        sheet.auto_filter.ref = f"A1:{get_column_letter(sheet.max_column)}{max(sheet.max_row, 1)}"
    for column in range(1, sheet.max_column + 1):
        values = [len(str(sheet.cell(row=row, column=column).value or "")) for row in range(1, min(sheet.max_row, 30) + 1)]
        column_letter = get_column_letter(column)
        sheet.column_dimensions[column_letter].width = min(max(max(values, default=10) + 2, 12), 42)
        if hidden_headers and sheet.cell(row=1, column=column).value in hidden_headers:
            sheet.column_dimensions[column_letter].hidden = True


def _implementation_values(entry: dict[str, Any] | None, key: str) -> list[str]:
    if not entry:
        return []
    return [str(implementation.get(key, "")) for implementation in entry.get("implementation", []) if implementation.get(key)]


def _execution_profiles(case: dict[str, Any], entry: dict[str, Any] | None) -> list[str]:
    profiles = {profile.get("name") for profile in case.get("execution", {}).get("profiles", []) if profile.get("name")}
    if entry:
        for implementation in entry.get("implementation", []):
            profiles.update(implementation.get("profiles", []))
    return sorted(profiles)


def _mock_services(case: dict[str, Any], entry: dict[str, Any] | None) -> list[str]:
    services = set(entry.get("required_mock_services", [])) if entry else set()
    for profile in case.get("execution", {}).get("profiles", []):
        if profile.get("name") == "mock":
            services.update(profile.get("required_services", []))
    return sorted(services)


def _manifest_validation(root: Path, stage: str, case: dict[str, Any], entry: dict[str, Any] | None) -> str:
    if not entry:
        if case.get("status") == "active" and case.get("automation") == "automated":
            return "未绑定"
        return "无需绑定"

    problems = []
    if entry.get("stage") != stage:
        problems.append("阶段不一致")
    if entry.get("contract_hash") != case_contract_hash(stage, case):
        problems.append("契约哈希失效")

    implementations = entry.get("implementation", [])
    files_exist = True
    for implementation in implementations:
        file_path = root / implementation.get("file", "")
        if not file_path.is_file():
            files_exist = False
            problems.append("脚本不存在")
            continue
        selector = implementation.get("selector", "")
        if selector and selector not in file_path.read_text(encoding="utf-8", errors="replace"):
            problems.append("Selector不存在")
    if implementations and files_exist and entry.get("implementation_hash") != implementation_hash(root, implementations):
        problems.append("实现哈希失效")
    return "；".join(dict.fromkeys(problems)) if problems else "校验通过"


def _case_rows(
    root: Path,
    cases: dict[str, tuple[str, dict[str, Any], Path]],
    manifest: dict[str, dict[str, Any]],
    stage: str,
):
    for case_id, (_, case, _) in sorted(cases.items()):
        if cases[case_id][0] != stage:
            continue
        stage_detail = {
            "D1": {"unit_boundary": case.get("unit_boundary"), "inputs": case.get("inputs"), "assertions": case.get("assertions")},
            "D2": case.get("contract"),
            "D3": {"integration_boundary": case.get("integration_boundary"), "observations": case.get("observations"), "cleanup": case.get("cleanup")},
            "D4": {"actor": case.get("actor"), "journey": case.get("journey"), "final_business_result": case.get("final_business_result"), "failure_evidence": case.get("failure_evidence")},
            "D5": {"risk_type": case.get("risk_type"), "risk": case.get("risk"), "condition": case.get("condition"), "metrics": case.get("metrics"), "thresholds": case.get("thresholds"), "recovery_expectations": case.get("recovery_expectations")},
        }[stage]
        entry = manifest.get(case_id)
        yield [
            case_id,
            case.get("feature_id"),
            case.get("requirement_ids"),
            case.get("business_rule_ids"),
            case.get("title"),
            case.get("type"),
            case.get("priority"),
            case.get("status"),
            case.get("automation"),
            case.get("objective"),
            case.get("preconditions"),
            case.get("test_data"),
            _steps(case),
            case.get("expected_results"),
            case.get("forbidden_side_effects"),
            stage_detail,
            case.get("execution"),
            entry.get("status") if entry else "unbound",
            _implementation_values(entry, "framework"),
            _implementation_values(entry, "file"),
            _implementation_values(entry, "selector"),
            _execution_profiles(case, entry),
            _mock_services(case, entry),
            entry.get("required_assets", []) if entry else [],
            _manifest_validation(root, stage, case, entry),
            entry.get("contract_hash", "") if entry else "",
            entry.get("implementation_hash", "") if entry else "",
        ]


def build_workbook(root: Path) -> Workbook:
    feature_documents = discover_documents(root, "features")
    case_documents = discover_documents(root, "cases")
    manifest_documents = discover_documents(root, "manifests")
    features, rules, _ = feature_index(feature_documents)
    cases, _ = case_index(case_documents)
    manifest, _ = manifest_index(manifest_documents)

    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.creator = "Nexent test asset generator"
    workbook.properties.lastModifiedBy = "Nexent test asset generator"
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    workbook.calculation.fullCalcOnLoad = False
    workbook.calculation.forceFullCalc = False

    info = workbook.create_sheet(SHEET_NAMES[0])
    info.sheet_view.showGridLines = False
    info.append(["Nexent formal test baseline"])
    info["A1"].fill = HEADER_FILL
    info["A1"].font = Font(color="FFFFFF", bold=True, size=14)
    info.append(["Source of truth", "Structured YAML/JSON under test/features, test/cases, test/changes and test/manifests"])
    info.append(["Generated view", "Seven read-only sheets; implementation and traceability fields are merged into each D1-D5 case row"])
    info.append(["Stages", "D1 unit/component; D2 API/contract; D3 integration/runtime; D4 browser journey; D5 risk validation"])
    info.append(["Legacy UT", "Kept separate and not counted as formal D1-D5 manifest coverage"])
    info.append(["Policy", "A2A is included across applicable stages; OAuth/CAS remains skipped by policy"])
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 110
    for row in info.iter_rows():
        for cell in row:
            cell.alignment = WRAP
    for cell in info[2]:
        cell.fill = SECTION_FILL

    feature_sheet = workbook.create_sheet(SHEET_NAMES[1])
    feature_rows = []
    for feature_id, feature in sorted(features.items()):
        owning_rules = [rule_id for rule_id, owner in rules.items() if owner == feature_id]
        feature_rows.append([feature_id, feature.get("name"), feature.get("status"), feature.get("description"), owning_rules, feature.get("required_test_stages"), feature.get("source")])
    _configure_sheet(feature_sheet, ["功能ID", "名称", "状态", "描述", "业务规则", "必需阶段", "来源"], feature_rows)

    case_headers = [
        "用例ID",
        "功能ID",
        "需求ID",
        "业务规则ID",
        "标题",
        "类型",
        "优先级",
        "状态",
        "自动化",
        "目标",
        "前置条件",
        "测试数据",
        "操作步骤及步骤预期",
        "最终预期",
        "禁止副作用",
        "阶段专属字段",
        "执行配置",
        "自动化实现状态",
        "测试框架",
        "脚本路径",
        "Selector",
        "执行Profile",
        "所需Mock服务",
        "所需测试资产",
        "Manifest校验状态",
        "契约哈希",
        "实现哈希",
    ]
    for stage, sheet_name in zip(("D1", "D2", "D3", "D4", "D5"), SHEET_NAMES[2:7]):
        _configure_sheet(
            workbook.create_sheet(sheet_name),
            case_headers,
            _case_rows(root, cases, manifest, stage),
            hidden_headers={"契约哈希", "实现哈希"},
        )
    return workbook


def workbook_snapshot(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, data_only=False)
    snapshot: dict[str, Any] = {"sheets": workbook.sheetnames, "data": {}}
    for sheet in workbook.worksheets:
        snapshot["data"][sheet.title] = {
            "freeze": str(sheet.freeze_panes or ""),
            "filter": sheet.auto_filter.ref,
            "gridlines": sheet.sheet_view.showGridLines,
            "columns": {
                key: {"width": value.width, "hidden": value.hidden}
                for key, value in sheet.column_dimensions.items()
            },
            "cells": [[cell.value for cell in row] for row in sheet.iter_rows()],
        }
    return snapshot


def write_workbook(root: Path, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(root).save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--check", action="store_true", help="Fail when the existing workbook differs from a freshly generated view")
    args = parser.parse_args()
    root = repository_root(args.root)
    output = (args.output or root / "test/generated/Nexent_测试基线.xlsx").resolve()
    if args.check:
        if not output.is_file():
            print(f"Generated Excel view is missing: {output}")
            return 1
        with tempfile.TemporaryDirectory() as temporary_directory:
            expected = Path(temporary_directory) / "expected.xlsx"
            write_workbook(root, expected)
            if workbook_snapshot(expected) != workbook_snapshot(output):
                print(f"Generated Excel view is stale or manually modified: {output}")
                return 1
        print(f"Generated Excel view is current: {output}")
        return 0
    write_workbook(root, output)
    print(f"Generated Excel view: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
